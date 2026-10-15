"""Import legacy time-series logs (txt/csv/xlsx) into SBL v1.

The importer guesses encoding, delimiter, header row, time format and a column
mapping (with confidence). Low-confidence columns become ``x.<name>`` so nothing
is lost; a human or the agent reviews the suggested mapping, saves it as JSON and
re-runs the conversion with ``--mapping``. Standard library only (xlsx needs
the optional ``openpyxl``).
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import io
import json
import math
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable

from . import sbl

ENCODINGS = ("utf-8-sig", "cp949", "euc-kr", "utf-16", "latin-1")

# Header aliases (lowercased, spaces/underscores/dashes removed) -> standard name.
# "{n}" marks a channel index captured from the header.
ALIASES: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"^(t|time|times|시간|경과시간|elapsed(time)?|testtime|sec|second)$"), "t"),
    (re.compile(r"^(timestamp|datetime|date|일시|날짜|absolutetime|realtime|시각)$"), "@datetime"),
    (re.compile(r"^(packv|packvolt(age)?|vpack|batteryvoltage|battv|totalvoltage|팩전압|전압|voltage|v)$"), "pack.v"),
    (re.compile(r"^(packi|packcurr(ent)?|ipack|batterycurrent|battcurrent|current|전류|i)$"), "pack.i"),
    (re.compile(r"^(packp|power|전력)$"), "pack.p"),
    (re.compile(r"^(cell|cellv|cellvolt(age)?|vcell|cv|셀전압|셀)_?(?P<n>\d+)$"), "cell.v.{n}"),
    (re.compile(r"^(temp|t|celltemp|tcell|ntc|온도|셀온도|temperature)_?(?P<n>\d+)$"), "temp.cell.{n}"),
    (re.compile(r"^(ambient|amb|tamb|ambienttemp|외기온도|주변온도)$"), "temp.amb"),
    (re.compile(r"^(soc|bmssoc|socbms|displaysoc|충전상태)$"), "bms.soc"),
    (re.compile(r"^(soh|bmssoh|sohbms|건강상태)$"), "bms.soh"),
    (re.compile(r"^(state|bmsstate|status|mode|상태)$"), "bms.state"),
    (re.compile(r"^(bal|balance|balancing|밸런싱)_?(?P<n>\d+)$"), "bal.{n}"),
    (re.compile(r"^(ahcap|ah|capacity|용량|chargeah)$"), "ref.ah"),
    (re.compile(r"^(stepindex|step|cycleindex|cycle|스텝|사이클)$"), "x.{raw}"),
]

UNIT_IN_HEADER = re.compile(r"^(?P<name>.*?)[\s_]*[\[(](?P<unit>[^\])]+)[\])]\s*$")
UNIT_ALIASES = {
    "v": "V", "mv": "mV", "a": "A", "ma": "mA", "w": "W", "kw": "kW", "s": "s", "sec": "s",
    "ms": "ms", "min": "min", "h": "h", "hr": "h", "°c": "degC", "℃": "degC", "degc": "degC",
    "c": "degC", "k": "K", "%": "%", "ah": "Ah", "mah": "mAh", "wh": "Wh", "kwh": "kWh",
    "ohm": "Ohm", "mohm": "mOhm", "mω": "mOhm", "ω": "Ohm", "mm": "mm", "um": "um",
    "kpa": "kPa", "pa": "Pa",
}
TIME_HMS = re.compile(r"^(?:(?P<d>\d+)[d.\s])?(?P<h>\d+):(?P<m>\d{1,2}):(?P<s>\d{1,2}(?:\.\d+)?)$")


@dataclass
class ColumnMap:
    source: str
    target: str
    unit: str
    confidence: float
    reason: str = ""


@dataclass
class ImportPlan:
    encoding: str
    delimiter: str
    header_row: int
    time_kind: str  # "seconds" | "hms" | "datetime" | "excel_serial" | "index"
    columns: list[ColumnMap] = field(default_factory=list)
    meta: dict[str, str] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, indent=2)

    @classmethod
    def from_json(cls, text: str) -> "ImportPlan":
        d = json.loads(text)
        d["columns"] = [ColumnMap(**c) for c in d.get("columns", [])]
        return cls(**d)


# ---------------------------------------------------------------- reading raw tables

def _read_rows(path: Path, encoding: str | None = None, delimiter: str | None = None):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as e:  # pragma: no cover
            raise RuntimeError("Excel import needs the optional 'openpyxl' package") from e
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["data"] if "data" in wb.sheetnames else wb.worksheets[0]
        rows = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows, "xlsx", "xlsx"
    raw = path.read_bytes()
    enc = encoding
    if enc is None:
        for cand in ENCODINGS:
            try:
                text = raw.decode(cand)
                enc = cand
                break
            except UnicodeDecodeError:
                continue
    text = raw.decode(enc or "latin-1")
    sample = "\n".join(text.splitlines()[:50])
    if delimiter is None:
        delimiter = _sniff_delimiter(sample)
    if delimiter == "ws":
        rows = [line.split() for line in text.splitlines()]
    else:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return rows, enc or "latin-1", delimiter


def _sniff_delimiter(sample: str) -> str:
    lines = [ln for ln in sample.splitlines() if ln.strip()][-20:]
    best, best_score = "ws", 0.0
    for d in ("\t", ",", ";", "|"):
        counts = [ln.count(d) for ln in lines]
        if not counts or max(counts) == 0:
            continue
        mode = max(set(counts), key=counts.count)
        score = counts.count(mode) / len(counts) * (1 if mode else 0) + mode * 1e-3
        if score > best_score:
            best, best_score = d, score
    return best


def _is_num(v) -> bool:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return True
    try:
        float(str(v).strip().replace(",", ""))
        return True
    except ValueError:
        return False


def _is_time_like(v) -> bool:
    if isinstance(v, (_dt.datetime, _dt.time)):
        return True
    s = str(v).strip()
    if TIME_HMS.match(s):
        return True
    try:
        _dt.datetime.fromisoformat(s.replace("/", "-").replace("Z", "+00:00"))
        return True
    except ValueError:
        return False


def _find_header(rows: list[list]) -> int:
    """Header = last mostly-text row right before a run of mostly-numeric rows."""
    def numeric_ratio(r):
        cells = [c for c in r if str(c).strip() != ""]
        if not cells:
            return 0.0
        return sum(1 for c in cells if _is_num(c) or _is_time_like(c)) / len(cells)

    for i in range(min(len(rows) - 1, 200)):
        r = rows[i]
        if len([c for c in r if str(c).strip()]) < 2:
            continue
        if numeric_ratio(r) < 0.5:
            nxt = rows[i + 1: i + 6]
            if nxt and all(numeric_ratio(x) >= 0.8 for x in nxt if any(str(c).strip() for c in x)):
                return i
    return 0


# ---------------------------------------------------------------- mapping inference

def _key(s: str) -> str:
    return re.sub(r"[\s_\-./#]+", "", s.strip().lower())


def _split_unit(header: str) -> tuple[str, str | None]:
    m = UNIT_IN_HEADER.match(header.strip())
    if m:
        u = UNIT_ALIASES.get(m.group("unit").strip().lower())
        if u:
            return m.group("name"), u
    return header, None


def _sanitize(name: str, fallback: str = "col") -> str:
    """ASCII channel-name segment; non-ASCII-only headers fall back to ``fallback``."""
    s = re.sub(r"[^a-z0-9_]+", "_", name.strip().lower()).strip("_")
    if not s:
        s = fallback
    if s[0].isdigit():
        s = "c" + s
    return s


def _meta_key(name: str) -> str:
    """Meta keys may keep Unicode letters (e.g. Korean labels from legacy headers)."""
    return re.sub(r"[^\w.]+", "_", name.strip().lower()).strip("_") or "key"


def _guess_unit(target: str, values: list[float]) -> tuple[str, str]:
    """Infer a unit from magnitudes when the header gave none."""
    finite = [abs(v) for v in values if v is not None and not math.isnan(v)]
    med = sorted(finite)[len(finite) // 2] if finite else 0.0
    if target.startswith("cell.v."):
        return ("mV", "median>100 → mV") if med > 100 else ("V", "range")
    if target.startswith("temp."):
        return ("K", "median>200 → K") if med > 200 else ("degC", "range")
    if target in ("bms.soc", "bms.soh"):
        return ("%", "range") if med > 1.5 else ("-", "0..1 fraction → scale x100")
    rule = sbl.standard_rule(target)
    return ((rule[0] if rule and rule[0] else "-"), "default")


def suggest_plan(path: str | Path, encoding: str | None = None) -> tuple[ImportPlan, list[list]]:
    p = Path(path)
    rows, enc, delim = _read_rows(p, encoding)
    h = _find_header(rows)
    header = [str(c).strip() for c in rows[h]]
    body = [r for r in rows[h + 1:] if any(str(c).strip() for c in r)]
    plan = ImportPlan(encoding=enc, delimiter=delim, header_row=h, time_kind="index")
    if h > 0:
        for r in rows[:h]:
            cells = [str(c).strip() for c in r if str(c).strip()]
            if len(cells) >= 2:
                plan.meta[f"legacy.{_meta_key(cells[0])}"] = " ".join(cells[1:])
    used: set[str] = set()
    time_col: str | None = None
    for j, col in enumerate(header):
        base, unit = _split_unit(col)
        k = _key(base)
        target, conf, reason = None, 0.0, ""
        for pat, tgt in ALIASES:
            m = pat.match(k)
            if m:
                n = m.groupdict().get("n")
                target = tgt.replace("{n}", f"{int(n):02d}" if n else "").replace("{raw}", _sanitize(base))
                conf, reason = (0.9 if n is None else 0.85), f"alias '{pat.pattern[:30]}'"
                break
        sample = [r[j] if j < len(r) else "" for r in body[:2000]]
        if target in ("t", "@datetime") or (target is None and j == 0 and sample and _is_time_like(sample[0])):
            if time_col is None:
                time_col = col
                plan.time_kind = _time_kind(sample, unit)
                target, conf = "t", max(conf, 0.8)
                reason = reason or "first column looks like time"
            else:
                target, conf, reason = f"x.{_sanitize(base, f'col{j:02d}')}", 0.5, "second time column"
        if target is None or target in used:
            target, conf, reason = f"x.{_sanitize(base, f'col{j:02d}')}", 0.3, reason or "no alias match"
        while target in used:  # keep channel names unique
            target += "_dup"
        if target != "t" and unit is None:
            nums = [_num(v) for v in sample]
            unit, why = _guess_unit(target, nums)
            reason += f"; unit {why}"
            conf = min(conf, 0.7) if why != "range" and why != "default" else conf
        if target == "t":
            unit = unit or "s"
        used.add(target)
        plan.columns.append(ColumnMap(col, target, unit or "-", round(conf, 2), reason))
    if time_col is None:
        plan.notes.append("no time column found: using row index as t (set sample period in meta)")
    names = {c.target for c in plan.columns}
    for need in ("pack.v", "pack.i"):
        if need not in names:
            plan.notes.append(f"'{need}' not identified: review mapping (L0 needs it)")
    plan.notes.append("current sign unknown: set meta current_sign after checking a charge segment")
    return plan, rows


def _time_kind(sample, unit) -> str:
    s = [v for v in sample if str(v).strip() != ""][:20]
    if not s:
        return "index"
    if all(isinstance(v, _dt.datetime) for v in s):
        return "datetime"
    if all(TIME_HMS.match(str(v).strip()) for v in s):
        return "hms"
    if all(_is_num(v) for v in s):
        vals = [float(str(v).replace(",", "")) for v in s]
        if unit is None and 20000 < vals[0] < 80000 and all(b >= a for a, b in zip(vals, vals[1:])):
            return "excel_serial"
        return "seconds"
    if all(_is_time_like(v) for v in s):
        return "datetime"
    return "index"


def _num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s == "" or s.lower() in ("nan", "na", "null", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return math.nan


def _time_value(v, kind: str, unit: str | None) -> float | None:
    if kind == "seconds":
        x = _num(v)
        if x is None or math.isnan(x):
            return x
        return x * sbl.UNITS.get(unit or "s", ("time", 1.0))[1]
    if kind == "hms":
        m = TIME_HMS.match(str(v).strip())
        if not m:
            return math.nan
        return (int(m.group("d") or 0) * 86400 + int(m.group("h")) * 3600
                + int(m.group("m")) * 60 + float(m.group("s")))
    if kind == "excel_serial":
        x = _num(v)
        return None if x is None else x * 86400.0
    if kind == "datetime":
        if isinstance(v, _dt.datetime):
            return v.timestamp() if v.tzinfo else v.replace(tzinfo=_dt.timezone.utc).timestamp()
        try:
            d = _dt.datetime.fromisoformat(str(v).strip().replace("/", "-").replace("Z", "+00:00"))
        except ValueError:
            return math.nan
        return d.timestamp() if d.tzinfo else d.replace(tzinfo=_dt.timezone.utc).timestamp()
    return None


# ---------------------------------------------------------------- conversion

def convert(path: str | Path, plan: ImportPlan | None = None, *, current_sign: str | None = None,
            sample_period_s: float = 1.0) -> tuple[sbl.SBLLog, ImportPlan]:
    p = Path(path)
    if plan is None:
        plan, rows = suggest_plan(p)
    else:
        rows, _, _ = _read_rows(p, plan.encoding, None if plan.delimiter == "xlsx" else plan.delimiter)
    header = [str(c).strip() for c in rows[plan.header_row]]
    body = [r for r in rows[plan.header_row + 1:] if any(str(c).strip() for c in r)]
    log = sbl.SBLLog(meta={"sbl": sbl.SPEC_VERSION})
    log.meta["source"] = p.name
    log.meta.update(plan.meta)
    if current_sign:
        log.meta["current_sign"] = current_sign

    idx = {h: j for j, h in enumerate(header)}
    tcol = next((c for c in plan.columns if c.target == "t"), None)
    if tcol is not None:
        j = idx[tcol.source]
        traw = [_time_value(r[j] if j < len(r) else "", plan.time_kind, tcol.unit) for r in body]
        t0 = next((x for x in traw if x is not None and not math.isnan(x)), 0.0)
        t = [None if x is None else x - t0 for x in traw]
        if plan.time_kind == "datetime" and "t0" not in log.meta:
            log.meta["t0"] = _dt.datetime.fromtimestamp(t0, _dt.timezone.utc).isoformat()
    else:
        t = [i * sample_period_s for i in range(len(body))]
        log.meta.setdefault("sampling", f"fixed:{sample_period_s}")
    log.channels.append(sbl.Channel("t", "s", "elapsed time", tcol.source if tcol else "<row index>"))
    log.data["t"] = t

    for c in plan.columns:
        if c.target == "t":
            continue
        j = idx.get(c.source)
        if j is None:
            continue
        vals = [_num(r[j] if j < len(r) else "") for r in body]
        unit = c.unit
        if unit == "-" and c.target in ("bms.soc", "bms.soh"):
            vals = [None if v is None else v * 100.0 for v in vals]
            unit = "%"
        log.channels.append(sbl.Channel(c.target, unit if unit in sbl.UNITS else "-", "", c.source))
        log.data[c.target] = vals
    sbl.normalize_units(log)
    return log, plan


def _cli(argv: Iterable[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m bmskit.logs.legacy",
                                 description="legacy txt/csv/xlsx -> SBL v1")
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("suggest", help="print a suggested import plan (JSON) for review")
    s.add_argument("path")
    s.add_argument("--encoding")
    c = sub.add_parser("convert", help="convert to SBL (.sbl.tsv or .xlsx)")
    c.add_argument("path")
    c.add_argument("dst")
    c.add_argument("--mapping", help="reviewed import plan JSON from 'suggest'")
    c.add_argument("--current-sign", choices=sbl.CURRENT_SIGNS)
    c.add_argument("--period", type=float, default=1.0, help="sample period if no time column")
    args = ap.parse_args(list(argv) if argv is not None else None)

    if args.cmd == "suggest":
        plan, _ = suggest_plan(args.path, args.encoding)
        print(plan.to_json())
        return 0
    plan = ImportPlan.from_json(Path(args.mapping).read_text(encoding="utf-8")) if args.mapping else None
    log, plan = convert(args.path, plan, current_sign=args.current_sign, sample_period_s=args.period)
    if args.dst.lower().endswith(".xlsx"):
        sbl.write_xlsx(log, args.dst)
    else:
        sbl.write_text(log, args.dst)
    rep = sbl.validate(log)
    print(f"wrote {args.dst}: conformance={rep.conformance} ok={rep.ok}")
    for i in rep.issues:
        if i.level != "info":
            print(i)
    return 0 if rep.ok else 1


if __name__ == "__main__":
    sys.exit(_cli())
