"""SBL (Standard BMS Log) v1 reference implementation.

Spec: docs/spec/sbl_v1.md. Standard library only; Excel (.xlsx) needs the optional
``openpyxl`` package.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

FORMAT_LINE = "#SBL 1.0"
SPEC_VERSION = "1.0"

NAME_RE = re.compile(r"^[a-z][a-z0-9_]*(\.[a-z0-9_]+)*$")
HEADER_RE = re.compile(r"^\s*(?P<name>[^\[\]]+?)\s*(\[(?P<unit>[^\]]*)\])?\s*$")
INDEX_RE = re.compile(r"^\d{2,}$")

CURRENT_SIGNS = ("discharge_positive", "charge_positive")
REQUIRED_META = ("current_sign",)
RECOMMENDED_META = (
    "t0", "source", "cell.chemistry", "cell.capacity_ah", "pack.series",
    "pack.parallel", "fw.version", "param_set", "test.type",
)

# Unit table: unit -> (dimension, factor to the canonical unit of that dimension).
UNITS: dict[str, tuple[str, float]] = {
    "s": ("time", 1.0), "ms": ("time", 1e-3), "min": ("time", 60.0), "h": ("time", 3600.0),
    "V": ("voltage", 1.0), "mV": ("voltage", 1e-3),
    "A": ("current", 1.0), "mA": ("current", 1e-3),
    "W": ("power", 1.0), "kW": ("power", 1e3),
    "degC": ("temperature", 1.0), "K": ("temperature_k", 1.0),
    "%": ("ratio_pct", 1.0),
    "Ah": ("charge", 1.0), "mAh": ("charge", 1e-3),
    "Wh": ("energy", 1.0), "kWh": ("energy", 1e3),
    "Ohm": ("resistance", 1.0), "mOhm": ("resistance", 1e-3),
    "mm": ("length", 1.0), "um": ("length", 1e-3),
    "kPa": ("pressure", 1.0), "Pa": ("pressure", 1e-3),
    "cycle": ("count", 1.0), "-": ("none", 1.0),
}
CANONICAL = {
    "time": "s", "voltage": "V", "current": "A", "power": "W", "temperature": "degC",
    "ratio_pct": "%", "charge": "Ah", "energy": "Wh", "resistance": "Ohm",
    "length": "mm", "pressure": "kPa", "count": "cycle", "none": "-",
}

# Standard channel families: (regex, canonical unit or None if free, level tag, plausible range)
STANDARD: list[tuple[re.Pattern[str], str | None, str, tuple[float, float] | None]] = [
    (re.compile(r"^t$"), "s", "L0", None),
    (re.compile(r"^pack\.v$"), "V", "L0", (0.0, 1500.0)),
    (re.compile(r"^pack\.i$"), "A", "L0", (-5000.0, 5000.0)),
    (re.compile(r"^pack\.p$"), "W", "L1", None),
    (re.compile(r"^cell\.v\.\d{2,}$"), "V", "L1", (0.0, 5.5)),
    (re.compile(r"^temp\.(cell|pack|module|busbar|fet|shunt)\.\d{2,}$"), "degC", "L1", (-60.0, 200.0)),
    (re.compile(r"^temp\.(amb|coolant_in|coolant_out|board)$"), "degC", "L1", (-60.0, 200.0)),
    (re.compile(r"^bal\.\d{2,}$"), "-", "L2", (0.0, 1.0)),
    (re.compile(r"^bms\.(soc|soh|soe|soc_raw|soc_min|soc_max)$"), "%", "L2", (-5.0, 105.0)),
    (re.compile(r"^bms\.(state|mode|fault_code)$"), "-", "L2", None),
    (re.compile(r"^bms\.(sop_chg|sop_dchg)$"), "W", "L2", None),
    (re.compile(r"^bms\.(sop_i_chg|sop_i_dchg)$"), "A", "L2", None),
    (re.compile(r"^bms\.ttf$"), "s", "L2", None),
    (re.compile(r"^bms\.capacity$"), "Ah", "L2", None),
    (re.compile(r"^bms\.isd_flag$"), "-", "L2", (0.0, 1.0)),
    (re.compile(r"^bms\.isd_cell$"), "-", "L2", None),
    (re.compile(r"^bms\.isd_r$"), "Ohm", "L2", None),
    (re.compile(r"^bms\.[a-z0-9_.]+$"), None, "L2", None),
    (re.compile(r"^flag\.[a-z0-9_]+$"), "-", "L2", (0.0, 1.0)),
    (re.compile(r"^model\.[a-z0-9_.]+$"), None, "L3", None),
    (re.compile(r"^ref\.v$"), "V", "R", None),
    (re.compile(r"^ref\.i$"), "A", "R", None),
    (re.compile(r"^ref\.soc$"), "%", "R", None),
    (re.compile(r"^ref\.ah$"), "Ah", "R", None),
    (re.compile(r"^ref\.soh$"), "%", "R", None),
    (re.compile(r"^ref\.[a-z0-9_.]+$"), None, "R", None),
    (re.compile(r"^chg\.(v_req)$"), "V", "L2", None),
    (re.compile(r"^chg\.(i_req)$"), "A", "L2", None),
    (re.compile(r"^chg\.[a-z0-9_.]+$"), None, "L2", None),
    (re.compile(r"^swell\.thickness(\.\d{2,})?$"), "mm", "S", None),
    (re.compile(r"^swell\.pressure(\.\d{2,})?$"), "kPa", "S", None),
    (re.compile(r"^x\.[a-z0-9_.]+$"), None, "X", None),
]


@dataclass
class Channel:
    name: str
    unit: str = "-"
    desc: str = ""
    source: str = ""


@dataclass
class SBLLog:
    """In-memory SBL log. ``data`` maps channel name to a list (None = missing)."""

    meta: dict[str, str] = field(default_factory=dict)
    channels: list[Channel] = field(default_factory=list)
    data: dict[str, list[float | None]] = field(default_factory=dict)

    @property
    def names(self) -> list[str]:
        return [c.name for c in self.channels]

    def __len__(self) -> int:
        return len(self.data.get("t", []))


@dataclass
class Issue:
    level: str  # "error" | "warning" | "info"
    code: str
    message: str

    def __str__(self) -> str:
        return f"[{self.level}] {self.code}: {self.message}"


@dataclass
class ValidationReport:
    issues: list[Issue] = field(default_factory=list)
    conformance: str = ""
    rows: int = 0
    channels: int = 0

    @property
    def ok(self) -> bool:
        return not any(i.level == "error" for i in self.issues)

    def add(self, level: str, code: str, message: str) -> None:
        self.issues.append(Issue(level, code, message))

    def to_dict(self) -> dict:
        return {
            "ok": self.ok, "conformance": self.conformance, "rows": self.rows,
            "channels": self.channels, "issues": [vars(i) for i in self.issues],
        }


def standard_rule(name: str):
    for pat, unit, tag, rng in STANDARD:
        if pat.match(name):
            return unit, tag, rng
    return None


def parse_header_cell(cell: str) -> tuple[str, str | None]:
    m = HEADER_RE.match(cell)
    if not m:
        return cell.strip(), None
    unit = m.group("unit")
    return m.group("name").strip(), (unit.strip() if unit is not None else None)


def _to_float(text: str) -> float | None:
    s = text.strip()
    if s == "" or s.lower() in ("nan", "na", "null", "none"):
        return None
    return float(s)


# ---------------------------------------------------------------- text I/O

def read_text(path: str | Path) -> SBLLog:
    """Read an SBL text file (.sbl.tsv / .sbl.csv). UTF-8, BOM allowed."""
    log = SBLLog()
    extra: dict[str, Channel] = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        first = f.readline().rstrip("\r\n")
        if not first.startswith("#SBL"):
            raise ValueError(f"{path}: first line must be '{FORMAT_LINE}', got {first!r}")
        log.meta["sbl"] = first[4:].strip()
        header: list[str] | None = None
        delim = "\t"
        for raw in f:
            line = raw.rstrip("\r\n")
            if header is None:
                if line.startswith("#@"):
                    key, _, value = line[2:].partition("=")
                    log.meta[key.strip().lower()] = value.strip()
                    continue
                if line.startswith("#!"):
                    ch = _parse_channel_line(line[2:])
                    extra[ch.name] = ch
                    continue
                if line.startswith("#") or not line.strip():
                    continue
                delim = "\t" if "\t" in line else ","
                header = line.split(delim)
                for cell in header:
                    name, unit = parse_header_cell(cell)
                    ch = extra.get(name, Channel(name))
                    if unit is not None:
                        ch.unit = unit
                    elif name in extra:
                        pass
                    else:
                        rule = standard_rule(name)
                        ch.unit = rule[0] if rule and rule[0] else "-"
                    log.channels.append(ch)
                    log.data[name] = []
                continue
            if not line.strip() or line.startswith("#"):
                continue
            cells = line.split(delim)
            if len(cells) < len(header):
                cells += [""] * (len(header) - len(cells))
            for ch, cell in zip(log.channels, cells):
                try:
                    log.data[ch.name].append(_to_float(cell))
                except ValueError:
                    log.data[ch.name].append(math.nan)  # flagged by validate()
    if header is None:
        raise ValueError(f"{path}: no header row")
    return log


def _parse_channel_line(text: str) -> Channel:
    fields = dict(
        (k.strip().lower(), v.strip())
        for k, _, v in (p.partition("=") for p in text.split(";")) if k.strip()
    )
    return Channel(
        name=fields.get("channel", ""), unit=fields.get("unit", "-"),
        desc=fields.get("desc", ""), source=fields.get("source", ""),
    )


def _fmt(v: float | None) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if float(v).is_integer() and abs(v) < 1e15:
        return str(int(v))
    return repr(float(v))


def write_text(log: SBLLog, path: str | Path, delimiter: str = "\t") -> None:
    """Write an SBL text file. Channel descriptions/sources go to '#!' lines."""
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(FORMAT_LINE + "\n")
        for k, v in log.meta.items():
            if k == "sbl":
                continue
            f.write(f"#@{k}={v}\n")
        for ch in log.channels:
            if ch.desc or ch.source:
                f.write(f"#!channel={ch.name};unit={ch.unit};desc={ch.desc};source={ch.source}\n")
        f.write(delimiter.join(f"{c.name}[{c.unit}]" for c in log.channels) + "\n")
        cols = [log.data[c.name] for c in log.channels]
        for row in zip(*cols):
            f.write(delimiter.join(_fmt(v) for v in row) + "\n")


# ---------------------------------------------------------------- Excel I/O

def read_xlsx(path: str | Path) -> SBLLog:
    """Read an SBL workbook: sheets 'meta' (key,value), optional 'channels', 'data'."""
    try:
        import openpyxl  # optional dependency (MIT)
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("Excel support needs the optional 'openpyxl' package") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    log = SBLLog()
    if "meta" in wb.sheetnames:
        for row in wb["meta"].iter_rows(values_only=True):
            if row and row[0] is not None:
                key = str(row[0]).strip()
                if key.lower() == "key":
                    continue
                log.meta[key.lower()] = "" if row[1] is None else str(row[1]).strip()
    extra: dict[str, Channel] = {}
    if "channels" in wb.sheetnames:
        for i, row in enumerate(wb["channels"].iter_rows(values_only=True)):
            if i == 0 or not row or row[0] is None:
                continue
            vals = [("" if v is None else str(v)) for v in row] + [""] * 4
            extra[vals[0]] = Channel(vals[0], vals[1] or "-", vals[2], vals[3])
    sheet = wb["data"] if "data" in wb.sheetnames else wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    for cell in header:
        name, unit = parse_header_cell(str(cell))
        ch = extra.get(name, Channel(name))
        if unit is not None:
            ch.unit = unit
        elif name not in extra:
            rule = standard_rule(name)
            ch.unit = rule[0] if rule and rule[0] else "-"
        log.channels.append(ch)
        log.data[name] = []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        for ch, v in zip(log.channels, row):
            if v is None or v == "":
                log.data[ch.name].append(None)
            else:
                try:
                    log.data[ch.name].append(float(v))
                except (TypeError, ValueError):
                    log.data[ch.name].append(math.nan)
    wb.close()
    log.meta.setdefault("sbl", SPEC_VERSION)
    return log


def write_xlsx(log: SBLLog, path: str | Path) -> None:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("Excel support needs the optional 'openpyxl' package") from e
    wb = openpyxl.Workbook()
    ws_meta = wb.active
    ws_meta.title = "meta"
    ws_meta.append(["key", "value"])
    ws_meta.append(["sbl", SPEC_VERSION])
    for k, v in log.meta.items():
        if k != "sbl":
            ws_meta.append([k, v])
    ws_ch = wb.create_sheet("channels")
    ws_ch.append(["name", "unit", "desc", "source"])
    for c in log.channels:
        ws_ch.append([c.name, c.unit, c.desc, c.source])
    ws = wb.create_sheet("data")
    ws.append([f"{c.name}[{c.unit}]" for c in log.channels])
    for row in zip(*(log.data[c.name] for c in log.channels)):
        ws.append([None if (v is None or (isinstance(v, float) and math.isnan(v))) else v for v in row])
    wb.save(path)


def read(path: str | Path) -> SBLLog:
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx(p)
    return read_text(p)


# ---------------------------------------------------------------- validation

def validate(log: SBLLog) -> ValidationReport:
    rep = ValidationReport(rows=len(log), channels=len(log.channels))
    version = log.meta.get("sbl", "")
    if not version.startswith("1."):
        rep.add("error", "E001", f"unsupported SBL version {version!r}")

    names = log.names
    if not names or names[0] != "t":
        rep.add("error", "E010", "first column must be 't' (seconds)")
    seen: set[str] = set()
    tags: set[str] = set()
    for ch in log.channels:
        if ch.name in seen:
            rep.add("error", "E011", f"duplicate channel {ch.name!r}")
        seen.add(ch.name)
        if not NAME_RE.match(ch.name):
            rep.add("error", "E012", f"invalid channel name {ch.name!r} (lowercase, dot-separated)")
            continue
        if ch.unit not in UNITS:
            rep.add("error", "E013", f"{ch.name}: unknown unit {ch.unit!r}")
            continue
        rule = standard_rule(ch.name)
        if rule is None:
            rep.add("warning", "W010", f"{ch.name}: not a standard name; use 'x.' prefix for custom channels")
            continue
        unit, tag, rng = rule
        tags.add(tag)
        if unit is not None and UNITS[ch.unit][0] != UNITS[unit][0]:
            rep.add("error", "E014", f"{ch.name}: unit {ch.unit!r} incompatible with {unit!r}")
        elif unit is not None and ch.unit != unit:
            rep.add("info", "I010", f"{ch.name}: non-canonical unit {ch.unit!r} (canonical {unit!r})")
        _check_values(rep, ch, log.data.get(ch.name, []), unit, rng)

    for idx_family in ("cell.v.", "bal."):
        idx = [n[len(idx_family):] for n in names if n.startswith(idx_family)]
        bad = [i for i in idx if not INDEX_RE.match(i)]
        if bad:
            rep.add("error", "E015", f"{idx_family}*: index must be zero-padded digits, got {bad[:3]}")

    for key in REQUIRED_META:
        if key not in log.meta:
            rep.add("error", "E020", f"missing required meta '{key}'")
    if log.meta.get("current_sign") not in (None, *CURRENT_SIGNS):
        rep.add("error", "E021", f"current_sign must be one of {CURRENT_SIGNS}")
    for key in RECOMMENDED_META:
        if key not in log.meta:
            rep.add("info", "I020", f"recommended meta '{key}' missing")
    t0 = log.meta.get("t0")
    if t0:
        try:
            parsed = _dt.datetime.fromisoformat(t0)
            if parsed.tzinfo is None:
                rep.add("warning", "W020", "t0 has no UTC offset")
        except ValueError:
            rep.add("error", "E022", f"t0 is not ISO 8601: {t0!r}")
    series = log.meta.get("pack.series")
    ncell = sum(1 for n in names if n.startswith("cell.v."))
    if series and ncell and series.isdigit() and int(series) != ncell:
        rep.add("warning", "W021", f"pack.series={series} but {ncell} cell.v channels")

    t = log.data.get("t", [])
    if t:
        back = sum(1 for a, b in zip(t, t[1:]) if a is not None and b is not None and b < a)
        dup = sum(1 for a, b in zip(t, t[1:]) if a is not None and a == b)
        if any(v is None for v in t):
            rep.add("error", "E030", "t has missing values")
        if back:
            rep.add("error", "E031", f"t goes backwards {back} time(s)")
        if dup:
            rep.add("warning", "W030", f"t has {dup} duplicate timestamp(s)")

    level = ""
    if {"t", "pack.v", "pack.i"} <= set(names):
        level = "L0"
        if any(n.startswith("cell.v.") for n in names) and any(n.startswith("temp.") for n in names):
            level = "L1"
            if any(n.startswith("bms.") for n in names):
                level = "L2"
                if "L3" in tags:
                    level = "L3"
    else:
        rep.add("error", "E040", "L0 requires t, pack.v, pack.i")
    suffix = "".join(f"+{s}" for s in ("R", "S") if s in tags)
    rep.conformance = (level or "none") + suffix
    return rep


def _check_values(rep, ch, values, unit, rng) -> None:
    n = len(values)
    if not n:
        return
    nan = sum(1 for v in values if v is not None and math.isnan(v))
    if nan:
        rep.add("error", "E032", f"{ch.name}: {nan} non-numeric value(s)")
    missing = sum(1 for v in values if v is None)
    if missing / n > 0.2:
        rep.add("warning", "W031", f"{ch.name}: {missing / n:.0%} missing")
    if rng and unit:
        factor = UNITS[ch.unit][1] / UNITS[unit][1] if UNITS[ch.unit][0] == UNITS[unit][0] else 1.0
        lo, hi = rng
        out = sum(1 for v in values if v is not None and not math.isnan(v) and not lo <= v * factor <= hi)
        if out:
            rep.add("warning", "W032", f"{ch.name}: {out} value(s) outside plausible range [{lo}, {hi}] {unit}")


def normalize_units(log: SBLLog) -> SBLLog:
    """Convert standard channels to canonical units in place (e.g. mV -> V)."""
    for ch in log.channels:
        rule = standard_rule(ch.name)
        if not rule or not rule[0] or ch.unit not in UNITS:
            continue
        target = rule[0]
        if ch.unit == target or UNITS[ch.unit][0] != UNITS[target][0]:
            continue
        k = UNITS[ch.unit][1] / UNITS[target][1]
        log.data[ch.name] = [None if v is None else v * k for v in log.data[ch.name]]
        ch.unit = target
    return log


# ---------------------------------------------------------------- CLI

def _cli(argv: Iterable[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m bmskit.logs.sbl", description="SBL v1 tools")
    sub = ap.add_subparsers(dest="cmd", required=True)
    v = sub.add_parser("validate", help="validate an SBL file (.sbl.tsv/.sbl.csv/.xlsx)")
    v.add_argument("path")
    v.add_argument("--json", action="store_true")
    c = sub.add_parser("convert", help="convert between SBL text and Excel")
    c.add_argument("src")
    c.add_argument("dst")
    c.add_argument("--normalize", action="store_true", help="convert to canonical units")
    args = ap.parse_args(list(argv) if argv is not None else None)

    if args.cmd == "validate":
        rep = validate(read(args.path))
        if args.json:
            print(json.dumps(rep.to_dict(), ensure_ascii=False, indent=2))
        else:
            print(f"conformance={rep.conformance} rows={rep.rows} channels={rep.channels} ok={rep.ok}")
            for i in rep.issues:
                print(i)
        return 0 if rep.ok else 1
    log = read(args.src)
    if args.normalize:
        normalize_units(log)
    if Path(args.dst).suffix.lower() == ".xlsx":
        write_xlsx(log, args.dst)
    else:
        write_text(log, args.dst, "," if args.dst.lower().endswith(".csv") else "\t")
    return 0


if __name__ == "__main__":
    sys.exit(_cli())
